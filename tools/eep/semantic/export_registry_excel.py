"""Export semantic-role-registry.json to an Excel review workbook.

This tool creates a human-reviewable workbook from the aggregated semantic role
registry. It does not replace the source mapping workbook. The source of truth
remains eep/semantic-role-mapping.xlsx.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook  # type: ignore[import-untyped]
    from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]
    from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
except ModuleNotFoundError as exc:  # pragma: no cover - dependency hint for CLI users
    raise SystemExit(
        "openpyxl is required for semantic role registry Excel export. "
        "Install the optional excel dependencies, e.g. `pip install -e .[excel]`."
    ) from exc


DEFAULT_SOURCE = Path("eep/semantic-role-registry.json")
DEFAULT_TARGET = Path("eep/semantic-role-registry.xlsx")
DEFAULT_SHEET = "semantic-role-registry"

COLUMNS = [
    "Semantic Role",
    "Channel Types",
    "Capabilities",
    "Function Keys",
    "Normalized Functions",
    "Attributes",
    "Confidences",
    "Review Statuses",
    "Notes",
]


def _load_roles(source: Path) -> list[dict[str, Any]]:
    payload = json.loads(source.read_text(encoding="utf-8"))

    if isinstance(payload, list):
        return [entry for entry in payload if isinstance(entry, dict)]

    roles = payload.get("roles") if isinstance(payload, dict) else None
    if not isinstance(roles, list):
        raise SystemExit(f"Invalid semantic role registry JSON: {source}")

    return [role for role in roles if isinstance(role, dict)]


def _format_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return " | ".join(str(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def export_registry_excel(
    source: Path = DEFAULT_SOURCE,
    target: Path = DEFAULT_TARGET,
    sheet_name: str = DEFAULT_SHEET,
) -> None:
    """Create an Excel review workbook from the aggregated registry JSON."""

    roles = _load_roles(source)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")

    for column_index, header in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.fill = header_fill
        cell.font = header_font

    for row_index, role in enumerate(roles, start=2):
        row_values = [
            role.get("semantic_role", ""),
            role.get("channel_types", []),
            role.get("capabilities", []),
            role.get("function_keys", []),
            role.get("normalized_functions", []),
            role.get("attributes", {}),
            role.get("confidences", []),
            role.get("review_statuses", []),
            role.get("notes", []),
        ]
        for column_index, value in enumerate(row_values, start=1):
            sheet.cell(row=row_index, column=column_index, value=_format_value(value))

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for column_index, header in enumerate(COLUMNS, start=1):
        column_letter = get_column_letter(column_index)
        max_length = len(header)
        for cell in sheet[column_letter]:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        sheet.column_dimensions[column_letter].width = min(max_length + 2, 100)

    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--target", type=Path, default=DEFAULT_TARGET)
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    args = parser.parse_args()

    export_registry_excel(args.source, args.target, args.sheet)
    print(f"Wrote {args.target}")


if __name__ == "__main__":
    main()
