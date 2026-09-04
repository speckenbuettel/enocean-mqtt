"""Import semantic-role-mapping.json back into an Excel workbook.

This reverse generator is intended for bulk human review. It recreates the
human-editable mapping table from the machine-readable JSON export while
preserving the canonical column order used by the source workbook.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook  # type: ignore[import-untyped]
    from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]
    from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
except ModuleNotFoundError as exc:  # pragma: no cover - dependency hint for CLI users
    raise SystemExit(
        "openpyxl is required for semantic role mapping import. "
        "Install the optional excel dependencies, e.g. `pip install -e .[excel]`."
    ) from exc


DEFAULT_SOURCE = Path("eep/semantic-role-mapping.json")
DEFAULT_TARGET = Path("eep/semantic-role-mapping.roundtrip.xlsx")
DEFAULT_SHEET = "semantic-role-mapping"

COLUMNS = [
    ("Function Key", "function_key"),
    ("Normalized Function", "normalized_function"),
    ("Semantic Role", "semantic_role"),
    ("Capability", "capability"),
    ("Confidence", "confidence"),
    ("Channel Type", "channel_type"),
    ("Attributes", "attributes"),
    ("Notes", "notes"),
    ("Review Status", "review_status"),
]


def _load_entries(source: Path) -> list[dict[str, Any]]:
    payload = json.loads(source.read_text(encoding="utf-8"))

    if isinstance(payload, list):
        # Compatibility with early prototype exports.
        return [entry for entry in payload if isinstance(entry, dict)]

    entries = payload.get("entries") if isinstance(payload, dict) else None
    if not isinstance(entries, list):
        raise SystemExit(f"Invalid semantic role mapping JSON: {source}")

    return [entry for entry in entries if isinstance(entry, dict)]


def _cell_value(entry: dict[str, Any], key: str) -> str:
    if key == "attributes":
        attributes_raw = entry.get("attributes_raw")
        if isinstance(attributes_raw, str):
            return attributes_raw

        attributes = entry.get("attributes")
        if isinstance(attributes, dict) and attributes:
            return json.dumps(attributes, ensure_ascii=False, sort_keys=True)
        return ""

    value = entry.get(key, "")
    if value is None:
        return ""
    return str(value)


def import_mapping(
    source: Path = DEFAULT_SOURCE,
    target: Path = DEFAULT_TARGET,
    sheet_name: str = DEFAULT_SHEET,
) -> None:
    """Recreate the review workbook from the mapping JSON."""

    entries = _load_entries(source)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")

    for column_index, (header, _) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.fill = header_fill
        cell.font = header_font

    for row_index, entry in enumerate(entries, start=2):
        for column_index, (_, key) in enumerate(COLUMNS, start=1):
            sheet.cell(row=row_index, column=column_index, value=_cell_value(entry, key))

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for column_index, (header, _) in enumerate(COLUMNS, start=1):
        column_letter = get_column_letter(column_index)
        max_length = len(header)
        for cell in sheet[column_letter]:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        sheet.column_dimensions[column_letter].width = min(max_length + 2, 80)

    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--target", type=Path, default=DEFAULT_TARGET)
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    args = parser.parse_args()

    import_mapping(args.source, args.target, args.sheet)
    print(f"Wrote {args.target}")


if __name__ == "__main__":
    main()
