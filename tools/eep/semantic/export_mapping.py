"""Export the semantic role mapping workbook to JSON.

The Excel workbook remains the human-editable source. This module exports the
mapping sheet to a stable, machine-readable JSON representation without
discarding review metadata.
"""

from __future__ import annotations

import argparse
import json
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook  # type: ignore[import-untyped]
except ModuleNotFoundError as exc:  # pragma: no cover - dependency hint for CLI users
    raise SystemExit(
        "openpyxl is required for semantic role mapping export. "
        "Install the optional excel dependencies, e.g. `pip install -e .[excel]`."
    ) from exc


DEFAULT_SOURCE = Path("eep/semantic-role-mapping.xlsx")
DEFAULT_TARGET = Path("eep/semantic-role-mapping.json")
DEFAULT_SHEET = "semantic-role-mapping"

EXPECTED_COLUMNS = [
    "Function Key",
    "Normalized Function",
    "Semantic Role",
    "Capability",
    "Confidence",
    "Channel Type",
    "Attributes",
    "Notes",
    "Review Status",
]


def _normalise_header(value: object) -> str:
    return str(value or "").strip()


def _json_key(header: str) -> str:
    return header.strip().lower().replace(" ", "_")


def _cell_to_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_attributes(raw_value: str) -> dict[str, Any]:
    """Parse the Attributes column.

    Empty cells become an empty object. Valid JSON objects are preserved as
    objects. Invalid non-empty cells are retained under ``raw`` so no source
    information is lost.
    """

    if not raw_value:
        return {}

    try:
        parsed = json.loads(raw_value)
    except json.JSONDecodeError:
        return {"raw": raw_value}

    if isinstance(parsed, dict):
        return parsed

    return {"raw": raw_value}


def export_mapping(
    source: Path = DEFAULT_SOURCE,
    target: Path = DEFAULT_TARGET,
    sheet_name: str = DEFAULT_SHEET,
) -> dict[str, Any]:
    """Export the semantic role mapping workbook to JSON."""

    workbook = load_workbook(source, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise SystemExit(
            f"Worksheet {sheet_name!r} not found in {source}. "
            f"Available sheets: {', '.join(workbook.sheetnames)}"
        )

    sheet = workbook[sheet_name]
    header_row = [_normalise_header(cell.value) for cell in sheet[1]]
    header_index = {header: idx for idx, header in enumerate(header_row) if header}

    missing = [column for column in EXPECTED_COLUMNS if column not in header_index]
    if missing:
        raise SystemExit(
            f"Missing required columns in {source}:{sheet_name}: {', '.join(missing)}"
        )

    entries: list[dict[str, Any]] = []

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        row_values = {
            column: _cell_to_text(row[header_index[column]]) for column in EXPECTED_COLUMNS
        }

        if not any(row_values.values()):
            continue

        entry = {
            "source_row": row_number,
            **{_json_key(column): row_values[column] for column in EXPECTED_COLUMNS},
        }

        attributes_raw = row_values["Attributes"]
        entry["attributes"] = _parse_attributes(attributes_raw)
        entry["attributes_raw"] = attributes_raw

        entries.append(entry)

    payload: dict[str, Any] = {
        "schema_version": "0.1",
        "generated_at": datetime.now(UTC).isoformat(timespec="seconds"),
        "source": str(source),
        "sheet": sheet_name,
        "columns": [_json_key(column) for column in EXPECTED_COLUMNS],
        "entries": entries,
    }

    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return payload


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--target", type=Path, default=DEFAULT_TARGET)
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    args = parser.parse_args()

    payload = export_mapping(args.source, args.target, args.sheet)
    print(f"Wrote {args.target} ({len(payload['entries'])} entries)")


if __name__ == "__main__":
    main()
